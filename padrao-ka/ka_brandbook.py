#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BRAND BOOK — motor de layout da terceira família de documentos do KA.

A Revelação de Essência e a Base Estratégica são apresentações: painel claro
sobre textura, uma ideia por página, muito ar. O **brand book** é outra
coisa — é manual. Tem cabeçalho com paginação, capítulos numerados, página
inteira de cor, e a informação organizada em duas colunas: o título à
esquerda, o conteúdo à direita.

A diagramação é a do Brand Book KA 2026, medida página a página no arquivo
da Kelly. O que muda de um cliente para outro é só a **paleta** e a
**marca** — a grade, a tipografia e os arquétipos ficam.

    from ka_brandbook import Brandbook, PALETA_KA
    d = Brandbook(marca="LUCAS MARTINI", paleta=PALETA_LM, logos={...})
    d.montar(PAGINAS)
    d.salvar("brandbook.pptx")

Tipografia: Outfit, como em todo o padrão KA. O peso vem do nome da fonte.
"""
import os
import sys

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from ka_layout import (  # noqa: E402
    ALT, LARG, BOLD, LIGHT, MEDIO, REG, SEMI, ULTRA, Deck,
)
from ka_paginas import altura, largura_texto, linhas  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(AQUI, "assets")
TRAMA_CANTO = os.path.join(ASSETS, "trama-canto.png")
TRAMA_DIREITA = os.path.join(ASSETS, "trama-direita.png")
TRAMA_PE = os.path.join(ASSETS, "trama-pe.png")

# ------------------------------------------------------------------- grade
ML, MR = 1.00, 19.00          # margens
W = MR - ML
COL_E, LARG_E = 1.00, 7.60    # coluna da esquerda: o título e o resumo
COL_D, LARG_D = 10.85, 7.50   # coluna da direita: o conteúdo
Y_CABECA = 0.68               # linha do cabeçalho
Y_TITULO = 1.90               # primeira linha do título
Y_CORPO = 2.05                # primeira linha da coluna da direita
FUNDO = 10.30                 # nada de texto abaixo disto

# -------------------------------------------------------------- tipografia
T_TITULO = 45       # título de página, caixa alta e peso leve
T_OLHO = 12.5       # eyebrow, caixa alta espacejada
T_CABECA = 11.5     # cabeçalho e paginação
T_LEAD = 21         # linha de abertura da coluna esquerda
T_CORPO = 19        # corpo
T_MINI = 15.5       # nota miúda
T_NOME = 25         # nome de coluna, de pilar, de versão
T_NUMERO = 44       # numeral de pilar
T_FRASE = 52        # frase grande numa página de silêncio
T_DIV_NUM = 58      # numeral do divisor de capítulo
T_DIV_NOME = 70     # nome do capítulo, em bold
T_CAPA = 21         # "BRAND BOOK" na capa
T_CAPA_MINI = 12.5  # edição, na capa


def _cor(hexa):
    return RGBColor.from_string(hexa.replace("#", "").upper())


# Paleta do Brand Book KA 2026 — medida no PDF dela.
PALETA_KA = {
    "escuro": "152535", "escuro2": "1F3449", "claro": "F8F7F2",
    "tinta": "152535", "corpo": "5A6066", "suave": "9AA0A6", "fio": "E2E0DA",
    "acento": "C87A44", "acento_escuro": "F0A337", "sobre_escuro": "F2F1EC",
    "corpo_escuro": "C9CFD6", "areia": "E8E4DA",
}


class Brandbook(Deck):
    """Manual de marca em 20 × 11,25 polegadas."""

    def __init__(self, marca, documento="BRAND BOOK", paleta=None, logos=None,
                 modelo=None, edicao=None):
        super().__init__(documento=documento, marca=marca, modelo=modelo)
        p = dict(PALETA_KA)
        p.update(paleta or {})
        self.p = {k: _cor(v) for k, v in p.items()}
        self.logos = logos or {}
        self.edicao = edicao
        self.n = 0              # página corrente
        self.total = 0          # preenchido em montar()
        self.avisos = []

    # ---------------------------------------------------------- utilidades
    def tinta(self, tom):
        """Cores de texto conforme o fundo da página."""
        if tom == "claro":
            return self.p["tinta"], self.p["corpo"], self.p["acento"], self.p["fio"]
        if tom == "areia":
            return self.p["tinta"], self.p["corpo"], self.p["acento"], self.p["fio"]
        return (self.p["sobre_escuro"], self.p["corpo_escuro"],
                self.p["acento_escuro"], self.p["escuro2"])

    def fundo(self, tom):
        return {"claro": self.p["claro"], "areia": self.p["areia"],
                "escuro2": self.p["escuro2"]}.get(tom, self.p["escuro"])

    def trama(self, s, caminho, opacidade=0.55):
        pic = s.shapes.add_picture(caminho, 0, 0, self.e(LARG), self.e(ALT))
        blip = pic._element.blipFill.find(qn("a:blip"))
        a = OxmlElement("a:alphaModFix")
        a.set("amt", str(int(opacidade * 100000)))
        blip.append(a)
        # a trama é fundo: vai para trás de tudo que já existe na página
        arvore = s.shapes._spTree
        arvore.remove(pic._element)
        arvore.insert(2, pic._element)
        return pic

    def marca_dagua(self, s, caminho, x, y, altura_alvo, opacidade=0.13):
        """Imagem grande e quase apagada, atras de tudo."""
        from PIL import Image as _Im
        iw, ih = _Im.open(caminho).size
        larg = altura_alvo * iw / ih
        pic = s.shapes.add_picture(caminho, self.e(x), self.e(y),
                                   self.e(larg), self.e(altura_alvo))
        blip = pic._element.blipFill.find(qn("a:blip"))
        a = OxmlElement("a:alphaModFix")
        a.set("amt", str(int(opacidade * 100000)))
        blip.append(a)
        arvore = s.shapes._spTree
        arvore.remove(pic._element)
        arvore.insert(2, pic._element)
        return pic

    def pagina(self, tom="claro", trama=None, cabecalho=True):
        s = self.prs.slides.add_slide(self.branco)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = self.fundo(tom)
        self.n += 1
        if trama:
            self.trama(s, trama, 0.5 if tom == "claro" else 0.55)
        if cabecalho:
            self.cabecalho(s, tom)
        return s

    def cabecalho(self, s, tom):
        """Filete de acento, nome do documento e a paginação. É ele que faz
        o documento parecer manual, e não apresentação."""
        _, _, acento, _ = self.tinta(tom)
        suave = self.p["suave"] if tom != "escuro" else self.p["corpo_escuro"]
        self.bloco(s, 0.98, Y_CABECA - 0.02, 0.022, 0.26, acento)
        self.texto(s, 1.12, Y_CABECA, 9.0,
                   "%s %s" % (self.documento, self.marca), T_CABECA,
                   suave, SEMI, spc=2.6, h=0.30)
        self.texto(s, MR - 4.0, Y_CABECA, 4.0,
                   "%02d / %02d" % (self.n, self.total or self.n), T_CABECA,
                   acento, SEMI, spc=2.6, h=0.30, align=PP_ALIGN.RIGHT)

    # ----------------------------------------------------------- tipografia
    def rico(self, p, texto, tam, cor, forte, fonte=LIGHT, fonte_forte=SEMI,
             spc=None):
        """Escreve um parágrafo com **negrito** e ~~riscado~~ no meio.

        No brand book o negrito carrega a frase que a pessoa precisa levar
        embora. Sem ele o texto vira um bloco cinza."""
        pedacos = []
        resto = texto
        while resto:
            for marca, tipo in (("**", "forte"), ("~~", "risco")):
                if resto.startswith(marca) and marca in resto[2:]:
                    fim = resto.index(marca, 2)
                    pedacos.append((tipo, resto[2:fim]))
                    resto = resto[fim + 2:]
                    break
            else:
                corte = len(resto)
                for marca in ("**", "~~"):
                    i = resto.find(marca)
                    if i > 0:
                        corte = min(corte, i)
                pedacos.append(("normal", resto[:corte]))
                resto = resto[corte:]
        for tipo, trecho in pedacos:
            if not trecho:
                continue
            r = self.txt(p, trecho, tam,
                         forte if tipo == "forte" else cor,
                         fonte_forte if tipo == "forte" else fonte, spc=spc)
            if tipo == "risco":
                r.font._rPr.set("strike", "sngStrike")

    def paragrafos(self, s, x, y, w, textos, tam, cor, forte, espaco=1.45,
                   gap=0.28, fonte=LIGHT):
        """Coluna de parágrafos medida linha a linha."""
        for t in textos:
            limpo = t.replace("**", "").replace("~~", "")
            h = altura(limpo, tam, w, espaco) + 0.06
            tf = self.caixa(s, x, y, w, h)
            self.rico(self.par(tf, True, espaco=espaco), t, tam, cor, forte,
                      fonte=fonte)
            y += h + gap
        return y

    def olho(self, s, texto, x, y, cor, w=7.0):
        self.texto(s, x, y, w, texto.upper(), T_OLHO, cor, SEMI, spc=2.8,
                   h=0.30)

    def titulo_pagina(self, s, titulo, tom, y=Y_TITULO, x=COL_E, fio_ate=None,
                      tam=T_TITULO):
        """Título em caixa alta, peso leve, e o filete que sai dele e corre
        para a direita — a assinatura da diagramação do brand book."""
        tinta, _, acento, fio = self.tinta(tom)
        if isinstance(titulo, str):
            titulo = [titulo]
        titulo = [t.upper() for t in titulo]
        n = len(titulo)
        self.texto(s, x, y, 9.2, titulo, tam, tinta, LIGHT, espaco=1.16,
                   spc=0.6, h=n * tam / 72.0 * 1.16 * 1.24 + 0.1)
        larg = max(largura_texto(t, tam) + n * 0.02 for t in titulo)
        y_fio = y + (n - 0.42) * tam / 72.0 * 1.16 * 1.24
        x_fio = x + larg + 0.55
        self.fio(s, x_fio, y_fio, (fio_ate or 8.90) - x_fio,
                 acento if fio_ate is None else fio, esp=0.014)
        return y + n * tam / 72.0 * 1.16 * 1.24 + 0.45

    # ------------------------------------------------------------ arquétipos
    def pg_capa(self, spec):
        s = self.pagina("escuro", TRAMA_DIREITA, cabecalho=False)
        self.trama(s, TRAMA_CANTO, 0.35)
        logo = self.logos.get("capa") or self.logos.get("escuro")
        if logo:
            lg = spec.get("logo_largura", 7.2)
            self.logo(s, logo, (LARG - lg) / 2, spec.get("logo_y", 3.30), lg)
        self.texto(s, 0, spec.get("y_titulo", 6.65), LARG,
                   spec.get("titulo", "BRAND BOOK").upper(), T_CAPA,
                   self.p["sobre_escuro"], LIGHT, spc=7.0,
                   align=PP_ALIGN.CENTER, h=0.46)
        if spec.get("edicao"):
            self.texto(s, 0, spec.get("y_titulo", 6.65) + 0.62, LARG,
                       spec["edicao"].upper(), T_CAPA_MINI,
                       self.p["acento_escuro"], SEMI, spc=4.0,
                       align=PP_ALIGN.CENTER, h=0.34)
        return s

    def pg_texto(self, spec):
        """A página mais comum: título à esquerda, conteúdo à direita."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        y = self.titulo_pagina(s, spec["titulo"], tom,
                               fio_ate=spec.get("fio_ate"))
        if spec.get("resumo"):
            self.paragrafos(s, COL_E, y, LARG_E - 0.6, spec["resumo"],
                            T_CORPO, corpo, tinta)
        if spec.get("marcado"):
            yy = y + (0.0 if not spec.get("resumo") else
                      sum(altura(t.replace("**", ""), T_CORPO, LARG_E - 0.6, 1.45)
                          + 0.34 for t in spec["resumo"]) + 0.30)
            self.texto(s, COL_E, yy, LARG_E - 0.6,
                       [l.upper() for l in spec["marcado"]], T_CORPO * 0.86,
                       tinta, SEMI, espaco=1.45, spc=1.6,
                       h=len(spec["marcado"]) * T_CORPO * 0.86 / 72.0 * 1.8)
        if spec.get("paragrafos"):
            self.paragrafos(s, COL_D, spec.get("y_direita", Y_CORPO), LARG_D,
                            spec["paragrafos"], T_CORPO, corpo, tinta)
        return s

    def pg_divisor(self, spec):
        """Virada de capítulo: página inteira de cor, numeral, nome em bold."""
        tom = spec.get("tom", "escuro")
        s = self.pagina(tom, TRAMA_DIREITA, cabecalho=True)
        self.trama(s, TRAMA_CANTO, 0.30)
        tinta, _, acento, _ = self.tinta(tom)
        if self.logos.get("escudo"):
            # a marca d'agua e sombra, nao ilustracao: entra quase apagada e
            # sangra pela direita, como as letras gigantes do brand book KA
            self.marca_dagua(s, self.logos["escudo"], 12.40, 0.60, 9.10,
                             spec.get("agua", 0.13))
        self.texto(s, ML, 4.55, 6.0, spec["numero"], T_DIV_NUM, acento, LIGHT,
                   h=T_DIV_NUM / 72.0 * 1.3)
        nome = spec["nome"]
        if isinstance(nome, str):
            nome = [nome]
        nome = [n.upper() for n in nome]
        self.texto(s, ML, 5.90, 12.0, nome, T_DIV_NOME, tinta, BOLD,
                   espaco=1.05, h=len(nome) * T_DIV_NOME / 72.0 * 1.05 * 1.24)
        y = 5.90 + len(nome) * T_DIV_NOME / 72.0 * 1.05 * 1.24 + 0.28
        self.fio(s, ML, y, 1.30, acento, esp=0.018)
        if spec.get("legenda"):
            self.texto(s, ML, y + 0.26, 12.0, spec["legenda"].upper(), T_OLHO,
                       tinta, SEMI, spc=3.0, h=0.32)
        return s

    def pg_lista(self, spec):
        """Linhas separadas por filete, com a seta de acento. Para listas de
        comportamento, de critério, de princípio."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        y = self.titulo_pagina(s, spec["titulo"], tom,
                               fio_ate=spec.get("fio_ate"))
        if spec.get("resumo"):
            self.paragrafos(s, COL_E, y, LARG_E - 0.6, spec["resumo"],
                            T_CORPO, corpo, tinta)
        yy = spec.get("y_direita", Y_CORPO)
        if spec.get("rotulo"):
            self.olho(s, spec["rotulo"], COL_D, yy, acento)
            yy += 0.42
        itens = spec["itens"]
        # mede a lista inteira antes de desenhar: com oito linhas e um fecho,
        # o respiro de 0,30" por item e o que joga o ultimo item para fora
        tam = T_CORPO
        for tam in (T_CORPO, T_CORPO * 0.92, T_MINI):
            alturas = [altura(i.replace("**", ""), tam, LARG_D - 0.55, 1.42)
                       for i in itens]
            fecho_h = (altura(spec["fecho"].replace("**", ""), tam, LARG_D,
                              1.45) + 0.40) if spec.get("fecho") else 0.0
            total = sum(alturas) + 0.52 * len(itens) + 0.22 + fecho_h
            if yy + total <= FUNDO:
                break
        folga = min(0.30, max(0.16,
                              (FUNDO - yy - total) / max(1, len(itens)) + 0.30))
        self.fio(s, COL_D, yy, LARG_D, fio, esp=0.008)
        yy += 0.22
        for item, h in zip(itens, alturas):
            self.texto(s, COL_D, yy - 0.015, 0.4, "→", tam * 0.82, acento,
                       LIGHT, h=0.3)
            tf = self.caixa(s, COL_D + 0.55, yy - 0.06, LARG_D - 0.55, h + 0.1)
            self.rico(self.par(tf, True, espaco=1.42), item, tam, corpo,
                      tinta)
            yy += h + folga
            self.fio(s, COL_D, yy, LARG_D, fio, esp=0.008)
            yy += 0.22
        if spec.get("fecho"):
            self.paragrafos(s, COL_D, yy + 0.20, LARG_D, [spec["fecho"]],
                            tam, corpo, tinta)
        return s

    def pg_colunas(self, spec):
        """Duas a quatro colunas de (rótulo, nome, texto). O filete fica
        entre o nome e o texto, nunca em volta."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        self.titulo_pagina(s, spec["titulo"], tom,
                           fio_ate=spec.get("fio_ate", MR))
        colunas = spec["colunas"]
        n = len(colunas)
        gap = 0.70
        larg = (W - gap * (n - 1)) / n
        y = spec.get("y", 4.60)
        # o filete fica na MESMA linha em todas as colunas: um nome que
        # quebra em duas linhas nao pode baixar so o seu filete
        alt_nome = max(linhas(nome.upper().replace("~~", ""), T_NOME, larg)
                       for _, nome, *_ in colunas) * T_NOME / 72.0 * 1.45
        fundo_col = y
        for i, col in enumerate(colunas):
            x = ML + i * (larg + gap)
            rot, nome = col[0], col[1]
            texto = col[2] if len(col) > 2 else ""
            if rot:
                self.olho(s, rot, x, y, acento, w=larg)
            nl = linhas(nome.upper().replace("~~", ""), T_NOME, larg)
            tf = self.caixa(s, x, y + 0.40, larg,
                            nl * T_NOME / 72.0 * 1.45 + 0.1)
            self.rico(self.par(tf, True, espaco=1.14), nome.upper(), T_NOME,
                      tinta, tinta, fonte=LIGHT, fonte_forte=SEMI, spc=0.4)
            ny = y + 0.40 + alt_nome
            self.fio(s, x, ny + 0.08, larg, acento, esp=0.012)
            if texto:
                fim = self.paragrafos(
                    s, x, ny + 0.36, larg,
                    texto if isinstance(texto, list) else [texto],
                    T_CORPO * 0.92, corpo, tinta, gap=0.22)
                fundo_col = max(fundo_col, fim)
        if spec.get("fecho"):
            self.paragrafos(s, ML, spec.get("y_fecho", fundo_col + 0.55),
                            W * 0.74, [spec["fecho"]], T_CORPO, corpo, tinta)
        return s

    def pg_numerais(self, spec):
        """Pilares: o numeral grande em acento, o nome em caixa alta, o
        filete e a leitura."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        self.titulo_pagina(s, spec["titulo"], tom,
                           fio_ate=spec.get("fio_ate", MR))
        itens = spec["itens"]
        n = len(itens)
        gap = 0.62
        larg = (W - gap * (n - 1)) / n
        y = spec.get("y", 4.70)
        for i, (nome, texto) in enumerate(itens):
            x = ML + i * (larg + gap)
            self.texto(s, x, y, larg, "%02d" % (i + 1), T_NUMERO, acento,
                       LIGHT, h=T_NUMERO / 72.0 * 1.30)
            ny = y + T_NUMERO / 72.0 * 1.30 + 0.06
            nlin = linhas(nome.upper(), T_OLHO * 1.12, larg)
            self.texto(s, x, ny, larg, nome.upper(), T_OLHO * 1.12, tinta,
                       SEMI, spc=2.2, espaco=1.45,
                       h=nlin * T_OLHO * 1.12 / 72.0 * 1.9)
            fy = ny + nlin * T_OLHO * 1.12 / 72.0 * 1.9 + 0.10
            self.fio(s, x, fy, larg, acento, esp=0.012)
            self.paragrafos(s, x, fy + 0.30, larg, [texto], T_CORPO * 0.92,
                            corpo, tinta, gap=0.20)
        return s

    def pg_declaracao(self, spec):
        """A página de silêncio do manual: uma frase, muito ar, e a cor."""
        tom = spec.get("tom", "escuro")
        s = self.pagina(tom, spec.get("trama", TRAMA_DIREITA))
        if spec.get("agua") and self.logos.get("escudo"):
            self.marca_dagua(s, self.logos["escudo"], 13.60, 1.60, 8.00, 0.10)
        tinta, corpo, acento, _ = self.tinta(tom)
        linhas_ = spec["linhas"]
        tam = spec.get("tam", T_FRASE)
        if spec.get("olho"):
            self.olho(s, spec["olho"], ML, spec.get("y", 5.30) - 0.60, acento)
        y = spec.get("y", 5.30)
        alt = len(linhas_) * tam / 72.0 * 1.14 * 1.24
        tf = self.caixa(s, ML, y, 16.0, alt + 0.2)
        for i, linha in enumerate(linhas_):
            p = self.par(tf, primeiro=(i == 0), espaco=1.14)
            self.rico(p, linha.upper(), tam, tinta, acento, fonte=LIGHT,
                      fonte_forte=BOLD)
        if spec.get("apoio"):
            self.paragrafos(s, ML, y + alt + 0.55, 12.0, spec["apoio"],
                            T_CORPO, corpo, tinta)
        return s

    def pg_pares(self, spec):
        """Em vez de / prefere. O que a marca risca e o que ela assina."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        y = self.titulo_pagina(s, spec["titulo"], tom,
                               fio_ate=spec.get("fio_ate"))
        if spec.get("resumo"):
            self.paragrafos(s, COL_E, y, LARG_E - 0.6, spec["resumo"],
                            T_CORPO, corpo, tinta)
        yy = spec.get("y_direita", Y_CORPO)
        rot_e = spec.get("rot_evitar", "em vez de")
        rot_p = spec.get("rot_preferir", "prefere")
        for evitar, preferir in spec["pares"]:
            self.fio(s, COL_D, yy, LARG_D, fio, esp=0.008)
            yy += 0.20
            self.olho(s, rot_e, COL_D, yy, self.p["suave"], w=LARG_D)
            h = altura(evitar, T_CORPO * 0.95, LARG_D, 1.42)
            tf = self.caixa(s, COL_D, yy + 0.34, LARG_D, h + 0.1)
            self.rico(self.par(tf, True, espaco=1.42), "~~%s~~" % evitar,
                      T_CORPO * 0.95, self.p["acento"], self.p["acento"])
            yy += 0.34 + h + 0.30
            self.olho(s, rot_p, COL_D, yy, acento, w=LARG_D)
            h2 = altura(preferir, T_CORPO * 0.95, LARG_D, 1.42)
            tf = self.caixa(s, COL_D, yy + 0.34, LARG_D, h2 + 0.1)
            self.rico(self.par(tf, True, espaco=1.42), preferir,
                      T_CORPO * 0.95, tinta, tinta)
            yy += 0.34 + h2 + 0.40
        return s

    def pg_meia(self, spec):
        """Meia página de cor: a frase de um lado, a imagem (ou o vazio) do
        outro. É a página de respiração do manual."""
        tom = spec.get("tom", "escuro")
        s = self.pagina(tom, None, cabecalho=False)
        corte = spec.get("corte", 9.40)
        lado = spec.get("lado", "esquerda")   # onde fica a imagem
        x_painel = corte if lado == "esquerda" else 0
        self.bloco(s, x_painel, 0, LARG - corte, ALT, self.fundo(tom))
        if spec.get("foto"):
            self.foto(s, 0 if lado == "esquerda" else LARG - corte, 0,
                      corte, ALT, spec["foto"], foco=spec.get("foco", 0.4))
        else:
            self.bloco(s, 0 if lado == "esquerda" else LARG - corte, 0, corte,
                       ALT, self.p["areia"])
            if self.logos.get("claro"):
                lg = 6.0
                self.logo(s, self.logos["claro"],
                          (corte - lg) / 2 if lado == "esquerda"
                          else LARG - corte + (corte - lg) / 2, 4.60, lg)
        tinta, corpo, acento, _ = self.tinta(tom)
        x = x_painel + 1.15
        larg = LARG - corte - 2.30
        self.cabecalho_meia(s, tom, x, larg)
        if spec.get("olho"):
            self.olho(s, spec["olho"], x, spec.get("y", 6.20) - 0.50, acento,
                      w=larg)
        y = spec.get("y", 6.20)
        tam = spec.get("tam", 32)
        tf = self.caixa(s, x, y, larg, len(spec["linhas"]) * tam / 72.0 * 1.9)
        for i, linha in enumerate(spec["linhas"]):
            p = self.par(tf, primeiro=(i == 0), espaco=1.22)
            self.rico(p, linha.upper(), tam, tinta, acento, fonte=LIGHT,
                      fonte_forte=BOLD)
        return s

    def cabecalho_meia(self, s, tom, x, larg):
        _, _, acento, _ = self.tinta(tom)
        suave = self.p["corpo_escuro"] if tom == "escuro" else self.p["suave"]
        self.bloco(s, x - 0.14, Y_CABECA - 0.02, 0.022, 0.26, acento)
        self.texto(s, x, Y_CABECA, larg * 0.7,
                   "%s %s" % (self.documento, self.marca), T_CABECA, suave,
                   SEMI, spc=2.6, h=0.30)
        self.texto(s, x, Y_CABECA, larg,
                   "%02d / %02d" % (self.n, self.total or self.n), T_CABECA,
                   acento, SEMI, spc=2.6, h=0.30, align=PP_ALIGN.RIGHT)

    def pg_paleta(self, spec):
        """As cores da marca, com o código embaixo — a página que o
        fornecedor abre primeiro."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        self.titulo_pagina(s, spec["titulo"], tom, fio_ate=MR)
        cores = spec["cores"]
        n = len(cores)
        gap = 0.55
        larg = (W - gap * (n - 1)) / n
        y = spec.get("y", 4.30)
        h = spec.get("altura", 3.10)
        for i, (nome, hexa, uso) in enumerate(cores):
            x = ML + i * (larg + gap)
            sp = self.bloco(s, x, y, larg, h, _cor(hexa))
            # amostra clara sobre papel claro some: ganha um fio de contorno
            if sum(_cor(hexa)) > 600:
                sp.line.color.rgb = self.p["fio"]
                sp.line.width = Pt(0.9)
            self.texto(s, x, y + h + 0.30, larg, nome.upper(), T_OLHO * 1.1,
                       tinta, SEMI, spc=2.2, h=0.34)
            self.texto(s, x, y + h + 0.68, larg, hexa.upper(), T_OLHO * 1.1,
                       acento, LIGHT, spc=2.2, h=0.34)
            self.paragrafos(s, x, y + h + 1.08, larg, [uso], T_MINI, corpo,
                            tinta, gap=0.18)
        return s

    def pg_marca(self, spec):
        """As versões do logo e o que não se faz com ele."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        self.titulo_pagina(s, spec["titulo"], tom, fio_ate=MR)
        versoes = spec["versoes"]
        n = len(versoes)
        gap = 0.60
        larg = (W - gap * (n - 1)) / n
        y = spec.get("y", 4.25)
        h = spec.get("altura", 3.00)
        for i, (nome, caminho, fundo_hex, legenda) in enumerate(versoes):
            x = ML + i * (larg + gap)
            self.bloco(s, x, y, larg, h, _cor(fundo_hex))
            lg = larg * 0.62
            from PIL import Image as _Im
            iw, ih = _Im.open(caminho).size
            alt_logo = lg * ih / iw
            if alt_logo > h * 0.62:
                alt_logo = h * 0.62
                lg = alt_logo * iw / ih
            self.logo(s, caminho, x + (larg - lg) / 2, y + (h - alt_logo) / 2,
                      lg)
            self.texto(s, x, y + h + 0.28, larg, nome.upper(), T_OLHO * 1.1,
                       tinta, SEMI, spc=2.2, h=0.34)
            self.paragrafos(s, x, y + h + 0.66, larg, [legenda], T_MINI, corpo,
                            tinta, gap=0.18)
        return s

    def pg_tipografia(self, spec):
        """O alfabeto da marca, em tamanho que dá para ver."""
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        if spec.get("olho"):
            self.olho(s, spec["olho"], COL_E, Y_TITULO - 0.52, acento)
        self.titulo_pagina(s, spec["titulo"], tom, fio_ate=MR)
        y = spec.get("y", 4.20)
        for nome, fonte, amostra, uso in spec["familias"]:
            self.olho(s, nome, ML, y, acento, w=W)
            self.texto(s, ML, y + 0.38, 12.6, amostra, spec.get("tam", 46),
                       tinta, fonte, espaco=1.1, h=0.95)
            self.paragrafos(s, 13.60, y + 0.46, 5.40, [uso], T_MINI, corpo,
                            tinta, gap=0.18)
            y += 1.70
            self.fio(s, ML, y, W, fio, esp=0.008)
            y += 0.52
        if spec.get("nota"):
            self.paragrafos(s, ML, y + 0.20, W * 0.72, [spec["nota"]], T_MINI,
                            self.p["suave"], corpo)
        return s

    def pg_fecho(self, spec):
        s = self.pagina(spec.get("tom", "escuro"), TRAMA_PE, cabecalho=False)
        logo = self.logos.get("escuro") or self.logos.get("capa")
        if logo:
            lg = spec.get("logo_largura", 6.4)
            self.logo(s, logo, (LARG - lg) / 2, spec.get("logo_y", 3.90), lg)
        if spec.get("frase"):
            self.texto(s, 0, spec.get("y", 6.90), LARG, spec["frase"].upper(),
                       T_OLHO * 1.25, self.p["acento_escuro"], SEMI, spc=4.2,
                       align=PP_ALIGN.CENTER, h=0.42)
        if spec.get("assinatura"):
            self.texto(s, 0, spec.get("y", 6.90) + 0.70, LARG,
                       spec["assinatura"], T_MINI, self.p["corpo_escuro"],
                       LIGHT, align=PP_ALIGN.CENTER, h=0.40)
        return s

    def pg_indice(self, spec):
        tom = spec.get("tom", "claro")
        s = self.pagina(tom, spec.get("trama"))
        tinta, corpo, acento, fio = self.tinta(tom)
        self.titulo_pagina(s, spec.get("titulo", "Conteúdo"), tom,
                           fio_ate=spec.get("fio_ate"))
        y = spec.get("y", 4.05)
        passo = spec.get("passo", 1.28)
        for num, nome, legenda in spec["itens"]:
            self.fio(s, COL_D - 0.0, y - 0.26, LARG_D + 0.0, fio, esp=0.008)
            self.texto(s, COL_E, y, 1.2, num, T_OLHO * 1.15, acento, SEMI,
                       spc=2.4, h=0.34)
            self.texto(s, COL_E + 1.4, y - 0.12, 7.0, nome.upper(), T_NOME,
                       tinta, LIGHT, spc=0.6, h=0.52)
            self.paragrafos(s, COL_D, y - 0.08, LARG_D, [legenda], T_CORPO,
                            corpo, tinta, gap=0.0)
            y += passo
        return s

    # -------------------------------------------------------------- montagem
    TIPOS = {
        "capa": "pg_capa", "indice": "pg_indice", "divisor": "pg_divisor",
        "texto": "pg_texto", "lista": "pg_lista", "colunas": "pg_colunas",
        "numerais": "pg_numerais", "declaracao": "pg_declaracao",
        "pares": "pg_pares", "meia": "pg_meia", "paleta": "pg_paleta",
        "marca": "pg_marca", "tipografia": "pg_tipografia", "fecho": "pg_fecho",
    }

    def montar(self, paginas):
        self.total = len(paginas)
        for spec in paginas:
            metodo = self.TIPOS.get(spec["tipo"])
            if not metodo:
                raise ValueError("tipo desconhecido: %s" % spec["tipo"])
            getattr(self, metodo)(spec)
        return self
