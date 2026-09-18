#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REVELACAO DE ESSENCIA - Flavia Pereira Muccelin
Remontagem integral do conteudo no padrao MODELO YUFIL / Base Estrategica
(ver MODELO-APRESENTACAO-YUFIL.md). Conteudo transcrito fielmente do original.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

CREME=RGBColor(0xF6,0xF5,0xF0); PRETO=RGBColor(0x1A,0x1A,0x1A); PRETO0=RGBColor(0,0,0)
CORPO=RGBColor(0x7A,0x7A,0x72); AREIA=RGBColor(0xC8,0xBF,0xA8); BEGE=RGBColor(0xE8,0xE4,0xD8)
SALVIA=RGBColor(0xD4,0xD9,0xCB); BEGE2=RGBColor(0xED,0xED,0xDD); BRANCO=RGBColor(0xFF,0xFF,0xFF)
F_TIT="Outfit"; F_BODY="Outfit"
EMU=914400; SW,SH=int(20*EMU),int(11.25*EMU)
ASSETS="/home/user/apresentacoes/assets"
TEXTURA=f"{ASSETS}/textura-papel.jpeg"; TRACO=f"{ASSETS}/traco.png"

prs=Presentation(); prs.slide_width=SW; prs.slide_height=SH
BLANK=prs.slide_layouts[6]

def add_slide(grey=False):
    s=prs.slides.add_slide(BLANK)
    bg=s.background; bg.fill.solid(); bg.fill.fore_color.rgb=CREME
    if grey: s.shapes.add_picture(TEXTURA,0,0,SW,SH)
    return s

def box(s,x,y,w,h):
    tb=s.shapes.add_textbox(Emu(int(x*EMU)),Emu(int(y*EMU)),Emu(int(w*EMU)),Emu(int(h*EMU)))
    tb.text_frame.word_wrap=True; return tb,tb.text_frame

def run(p,text,size,color,font=F_BODY,bold=False,italic=False,spc=None):
    r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.name=font
    r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=color
    if spc is not None: r._r.get_or_add_rPr().set('spc',str(int(spc*100)))
    return r

def rect(s,x,y,w,h,color):
    sp=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Emu(int(x*EMU)),Emu(int(y*EMU)),Emu(int(w*EMU)),Emu(int(h*EMU)))
    sp.fill.solid(); sp.fill.fore_color.rgb=color; sp.line.fill.background(); sp.shadow.inherit=False; return sp

def eyebrow(s,text,x=1.7,y=0.9,w=15.0,color=CORPO):
    tb,tf=box(s,x,y,w,0.5); run(tf.paragraphs[0],text.upper(),16.5,color,F_BODY,spc=2.5)

def h1(s,text,x=1.7,y=1.55,size=39,w=16.6):
    tb,tf=box(s,x,y,w,1.4); run(tf.paragraphs[0],text,size,PRETO,F_BODY,bold=True)

def txt(s,text,x,y,w,h,size,color=CORPO,bold=False,italic=False,font=F_BODY):
    tb,tf=box(s,x,y,w,h); run(tf.paragraphs[0],text,size,color,font,bold=bold,italic=italic); return tf

def lista(s,itens,x,y,w,h,size=23,color=CORPO,marker=None,mcolor=None,bold=False,space=None):
    tb,tf=box(s,x,y,w,h); first=True
    for t in itens:
        p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
        if space is not None: p.space_after=Pt(space)
        if marker: run(p,marker+"  ",size,mcolor or AREIA,F_BODY,bold=True)
        run(p,t,size,color,F_BODY,bold=bold)
    return tf

def divisor(s,num,titulo):
    """Divisor de grande tema: fundo cinza + numero areia + titulo 100pt."""
    tb,tf=box(s,2.2,3.55,6.0,0.7)
    run(tf.paragraphs[0],num,21,AREIA,F_BODY,bold=True,spc=3)
    tb,tf=box(s,2.2,4.25,16.6,3.0)
    run(tf.paragraphs[0],titulo,100,PRETO0,F_BODY,bold=True)

def nota_rodape(s,text,y=10.05,size=20):
    txt(s,text,1.7,y,16.6,0.9,size,CORPO)

# =====================================================================
# 01 · CAPA  (cinza)
# =====================================================================
s=add_slide(grey=True)
tb,tf=box(s,1.68,3.93,14.9,2.5)
run(tf.paragraphs[0],"REVELAÇÃO DE",75,PRETO0,F_TIT,bold=True)
run(tf.add_paragraph(),"ESSÊNCIA",75,PRETO0,F_TIT,bold=True,spc=1)
s.shapes.add_picture(TRACO,Emu(int(13.9*EMU)),Emu(int(4.13*EMU)),Emu(int(1.63*EMU)),Emu(int(1.94*EMU)))
tb,tf=box(s,1.68,6.45,14.5,0.9); p=tf.paragraphs[0]
run(p,"Flávia Pereira Muccelin",30,PRETO,F_BODY,bold=True)
run(p,"  |  Origem Identitária da Fundadora",30,PRETO,F_BODY)
tb,tf=box(s,1.68,8.52,9.0,1.2)
run(tf.paragraphs[0],"SETEMBRO/2026",19.5,PRETO0,F_BODY,spc=1.5)
p2=tf.add_paragraph()
run(p2,"Método ",26.5,PRETO,F_BODY); run(p2,"Marca com Essência ",26.5,PRETO,F_BODY,bold=True)
run(p2,"©",15,PRETO,F_BODY,bold=True)

# =====================================================================
# 02 · SUMÁRIO  (cinza)
# =====================================================================
s=add_slide(grey=True)
tb,tf=box(s,1.9,0.9,10,1.6); run(tf.paragraphs[0],"SUMÁRIO",78,PRETO0,F_TIT,bold=True)
sumario=[("01","A origem","As experiências que formaram a visão de mundo da Flávia."),
         ("02","As tensões","Os momentos que a colocaram em movimento."),
         ("03","Os padrões","A lógica que se repete ao longo da trajetória."),
         ("04","O Ikigai","O que gera sentido e realização para a Flávia."),
         ("05","A essência","O princípio que conecta a sua história."),
         ("06","As marcas","Como essa essência se manifesta na origem da Forma e da My Home.")]
ys=[2.9,4.55,6.2]
for i,(num,tit,sub) in enumerate(sumario):
    x = 1.9 if i<3 else 10.3
    y = ys[i%3]
    tb,tf=box(s,x,y,1.4,1.0); run(tf.paragraphs[0],num,43,PRETO,F_BODY,bold=True)
    tb,tf=box(s,x+1.45,y+0.05,6.5,1.5)
    run(tf.paragraphs[0],tit,30,PRETO,F_BODY,bold=True)
    run(tf.add_paragraph(),sub,20,CORPO,F_BODY)

# =====================================================================
# 03 · DIVISOR DE ABERTURA  (cinza)
# =====================================================================
s=add_slide(grey=True)
tb,tf=box(s,2.2,3.9,15.4,3.6); p=tf.paragraphs[0]
run(p,"POR QUE COMEÇAMOS",100,PRETO0,F_BODY)
p2=tf.add_paragraph(); run(p2,"PELA ",100,PRETO0,F_BODY); run(p2,"ESSÊNCIA?",100,PRETO0,F_BODY,bold=True)

# =====================================================================
# 04 · O QUE É / O QUE NÃO É
# =====================================================================
s=add_slide()
eyebrow(s,"Revelação de essência")
h1(s,"O PONTO DE PARTIDA DO MÉTODO")
txt(s,"A Forma e a My Home são negócios distintos, mas nasceram e se desenvolveram a partir da visão de uma mesma fundadora.",
    1.7,3.0,14.6,1.6,30,PRETO,bold=True)
txt(s,"Em marcas fortemente ligadas à fundadora, compreender a história de quem as criou ajuda a identificar crenças, princípios e formas de pensar e agir que influenciaram a construção dos negócios.",
    1.7,4.7,14.6,1.8,23,CORPO)
rect(s,1.7,6.9,7.9,2.9,BEGE)
eyebrow(s,"O que não é",2.15,7.25,6.8)
txt(s,"Uma análise psicológica da Flávia.",2.15,7.9,6.9,1.6,26,PRETO,bold=True)
rect(s,10.4,6.9,7.9,2.9,SALVIA)
eyebrow(s,"O que é",10.85,7.25,6.8)
txt(s,"A investigação da origem identitária das marcas a partir da fundadora.",10.85,7.9,6.9,1.8,26,PRETO,bold=True)

# =====================================================================
# 05 · A PERGUNTA QUE CONDUZ
# =====================================================================
s=add_slide()
eyebrow(s,"A pergunta que conduz esta etapa",2.3,3.0)
tb,tf=box(s,2.3,3.7,15.4,4.2); p=tf.paragraphs[0]
run(p,"O que já fazia parte de quem a Flávia ",62,PRETO,F_BODY)
run(p,"era antes",62,PRETO,F_BODY,bold=True)
run(p," da Forma e da My Home?",62,PRETO,F_BODY)
rect(s,2.5,8.3,15,0.04,AREIA)

# =====================================================================
# 06 · MÉTODO — COMO CHEGAMOS À REVELAÇÃO
# =====================================================================
s=add_slide()
eyebrow(s,"Método")
h1(s,"COMO CHEGAMOS À REVELAÇÃO")
txt(s,"A essência não é definida por uma característica isolada. Ela se revela quando diferentes momentos da trajetória apresentam uma mesma lógica.",
    1.7,2.9,14.6,1.4,23,CORPO)
metodo=[("História","Mostra as experiências e as referências que formaram a visão de mundo da Flávia."),
        ("Tensões","Mostram os momentos em que aquilo que existia entrou em conflito com aquilo que ela desejava ou conseguia enxergar."),
        ("Padrões","Revelam como a Flávia responde, repetidamente, a essas situações."),
        ("Ikigai","Revela o que gera sentido, realização e vontade de contribuir.")]
for i,(tit,desc) in enumerate(metodo):
    x=1.7+i*4.23
    rect(s,x,4.5,3.95,3.4,BRANCO if i%2==0 else BEGE)
    tb,tf=box(s,x+0.4,4.8,3.2,2.9)
    run(tf.paragraphs[0],f"{i+1:02d}",27,AREIA,F_BODY,bold=True)
    run(tf.add_paragraph(),tit,24,PRETO,F_BODY,bold=True)
    run(tf.add_paragraph(),desc,17,CORPO,F_BODY)
tb,tf=box(s,1.7,8.35,16.6,1.0); p=tf.paragraphs[0]
run(p,"História  +  tensões  +  padrões  +  Ikigai  =  ",26,CORPO,F_BODY)
run(p,"essência",26,PRETO,F_BODY,bold=True)
nota_rodape(s,"A essência é o princípio central que conecta essas dimensões e revela a lógica que orienta a forma como a Flávia enxerga, age e transforma.",9.55,18)

# =====================================================================
# 07 · DIVISOR 01 — A ORIGEM  (cinza)
# =====================================================================
s=add_slide(grey=True); divisor(s,"01","A ORIGEM")

# =====================================================================
# 08 · ONDE ESSA HISTÓRIA COMEÇA
# =====================================================================
s=add_slide()
rect(s,12.9,0,7.1,11.25,SALVIA)
eyebrow(s,"A origem",1.5,1.0,10.5)
h1(s,"ONDE ESSA HISTÓRIA COMEÇA",1.5,1.6,39,11.0)
txt(s,"A história familiar da Flávia é marcada por limitações, trabalho, fé e recomeços.",
    1.5,3.3,10.8,1.4,27,PRETO,bold=True)
txt(s,"A família deixou Guiratinga e foi para Primavera do Leste em busca de uma vida melhor. O pai da Flávia passou a trabalhar como caseiro justamente na propriedade onde, décadas depois, ela viveria como proprietária.",
    1.5,5.0,10.8,2.4,23,CORPO)
rect(s,1.5,7.6,10.8,1.5,BRANCO)
txt(s,"A Flávia chegou àquele lugar como “a filha do peão”.",1.9,7.9,10.0,1.0,26,PRETO,bold=True)
txt(s,"TRABALHO   ·   FAMÍLIA   ·   FÉ   ·   HONESTIDADE   ·   GRATIDÃO",1.5,9.5,11.0,0.6,18,PRETO,bold=True)

# =====================================================================
# 09 · OS PRIMEIROS SINAIS
# =====================================================================
s=add_slide()
eyebrow(s,"A origem")
h1(s,"OS PRIMEIROS SINAIS")
txt(s,"A vontade de construir uma realidade diferente aparece muito antes da criação das empresas.",
    1.7,3.0,8.6,1.6,27,PRETO,bold=True)
txt(s,"Ainda criança, a Flávia acompanhava o pai no garimpo e chegou a cozinhar para os trabalhadores. Mais tarde, ia de bicicleta da chácara até o centro para trabalhar. Também estudava e jogava futsal para conquistar uma bolsa que ajudasse a manter a faculdade.",
    1.7,4.9,8.6,3.2,23,CORPO)
txt(s,"Ela mesma diz:",10.9,2.9,7.4,0.6,18,CORPO)
rect(s,10.9,3.5,7.4,2.3,BRANCO)
txt(s,"“",11.2,3.35,1.5,1.4,56,AREIA,bold=True)
txt(s,"Eu não quero isso para mim.",11.9,4.2,6.1,1.4,26,PRETO,bold=True)
txt(s,"E, mais tarde, ao perceber que o cargo limitaria o seu crescimento:",10.9,6.2,7.4,0.9,18,CORPO)
rect(s,10.9,7.2,7.4,2.3,BEGE)
txt(s,"“",11.2,7.05,1.5,1.4,56,AREIA,bold=True)
txt(s,"Eu queria mais.",11.9,7.9,6.1,1.4,26,PRETO,bold=True)

# =====================================================================
# 10-11 · AS EXPERIÊNCIAS QUE A FORMARAM
# =====================================================================
EXPERIENCIAS=[[("01","A escassez","A Flávia cresceu em uma realidade de poucos recursos.",
                "A vontade de ampliar possibilidades e não aceitar a condição presente como limite."),
               ("02","O portão","Ainda adolescente, a Flávia entrava escondida na casa dos donos da fazenda onde o pai trabalhava e se imaginava vivendo aquela realidade. Anos depois, tornou-se proprietária daquele mesmo lugar.",
                "A capacidade de se enxergar dentro de uma realidade antes de ela existir concretamente.")],
              [("03","O trabalho e o estudo","Trabalho, faculdade, bicicleta e futsal como caminho para conquistar uma bolsa de estudos.",
                "Para a Flávia, enxergar uma possibilidade exige movimento para torná-la real."),
               ("04","A trajetória dos pais","A Flávia cresceu vendo os pais recomeçarem sem abandonar a honestidade, a fé, o trabalho e a família.",
                "A forma de construir importa tanto quanto aquilo que é conquistado.")]]
for pagina,grupo in enumerate(EXPERIENCIAS):
    s=add_slide()
    eyebrow(s,"A origem")
    h1(s,"AS EXPERIÊNCIAS QUE A FORMARAM" + ("" if pagina==0 else " ·  CONTINUAÇÃO"))
    for i,(num,tit,desc,revela) in enumerate(grupo):
        x=1.7+i*8.8
        rect(s,x,3.2,8.0,6.6,BRANCO if i==0 else BEGE)
        tb,tf=box(s,x+0.5,3.55,7.0,1.6)
        run(tf.paragraphs[0],num,27,AREIA,F_BODY,bold=True)
        run(tf.add_paragraph(),tit,30,PRETO,F_BODY,bold=True)
        txt(s,desc,x+0.5,5.3,7.0,2.4,20,CORPO)
        rect(s,x+0.5,7.75,7.0,0.03,AREIA)
        eyebrow(s,"O que isso revela",x+0.5,7.95,6.8)
        txt(s,revela,x+0.5,8.5,7.0,1.2,21,PRETO,bold=True)

# =====================================================================
# 12 · DIVISOR 02 — AS TENSÕES  (cinza)
# =====================================================================
s=add_slide(grey=True); divisor(s,"02","AS TENSÕES")

# =====================================================================
# 13 · AS TENSÕES QUE A COLOCARAM EM MOVIMENTO
# =====================================================================
s=add_slide()
eyebrow(s,"As tensões")
h1(s,"AS TENSÕES QUE A COLOCARAM EM MOVIMENTO")
txt(s,"Existe uma lógica recorrente na trajetória da Flávia.",1.7,2.75,14.6,0.7,23,CORPO)
tensoes=[("A realidade dizia","“Essa é a sua condição.”","A Flávia enxergava outra."),
         ("O cargo dizia","“Até aqui você pode chegar.”","A Flávia queria mais."),
         ("As marcas prontas diziam","“É assim que deve ser feito.”","A Flávia queria fazer do seu jeito."),
         ("A maternidade exigiu uma escolha","Conciliar a maternidade e o empreendedorismo não era possível como ela desejava naquele momento.","Ela parou e, anos depois, construiu um caminho de volta."),
         ("As sobras tinham um destino","“O descarte.”","A Flávia enxergou matéria para uma nova criação.")]
y0=3.55; hh=1.28; gap=0.13
for i,(rot,fala,resp) in enumerate(tensoes):
    y=y0+i*(hh+gap)
    rect(s,1.7,y,9.3,hh,BRANCO if i%2==0 else BEGE)
    eyebrow(s,rot,2.05,y+0.13,8.6)
    txt(s,fala,2.05,y+0.58,8.7,0.7,19,PRETO)
    txt(s,"→",11.25,y+0.3,0.8,0.7,24,AREIA,bold=True)
    txt(s,resp,12.1,y+0.35,6.2,0.9,21,PRETO,bold=True)

# =====================================================================
# 14 · DIVISOR 03 — OS PADRÕES  (cinza)
# =====================================================================
s=add_slide(grey=True); divisor(s,"03","OS PADRÕES")

# =====================================================================
# 15 · O PADRÃO INVISÍVEL
# =====================================================================
s=add_slide()
rect(s,0,0,7.6,11.25,BEGE2)
eyebrow(s,"Os padrões",1.3,2.9,5.8)
txt(s,"O PADRÃO",1.3,3.5,6.0,1.2,51,PRETO,bold=True)
txt(s,"INVISÍVEL",1.3,4.6,6.0,1.2,51,PRETO,bold=True)
txt(s,"As situações são diferentes. A resposta da Flávia segue a mesma lógica.",
    1.3,6.1,5.9,2.0,24,CORPO)
passos=["Ela encontra uma realidade ou um limite.",
        "Enxerga que aquilo pode ser diferente.",
        "Constrói um caminho para transformar essa possibilidade em realidade."]
for i,t in enumerate(passos):
    y=2.9+i*2.1
    rect(s,9.0,y,9.3,1.75,BRANCO)
    txt(s,f"{i+1:02d}",9.45,y+0.35,1.0,0.9,27,AREIA,bold=True)
    txt(s,t,10.6,y+0.4,7.3,1.1,23,PRETO,bold=True)
txt(s,"O padrão está na forma como a Flávia responde quando percebe que existe uma possibilidade além daquilo que está dado.",
    9.0,9.4,9.3,1.4,20,CORPO)

# =====================================================================
# 16 · O MOVIMENTO QUE SE REPETE
# =====================================================================
s=add_slide()
eyebrow(s,"Os padrões")
h1(s,"O MOVIMENTO QUE SE REPETE")
movimento=["Enxerga a realidade como ela é","Percebe que aquilo não precisa ser definitivo",
           "Enxerga outra possibilidade","Procura um caminho",
           "Busca conhecimento, mobiliza pessoas e recursos","Transforma a possibilidade em realidade"]
for i,t in enumerate(movimento):
    x=1.7+i*2.78
    rect(s,x,3.5,2.5,3.6,BRANCO if i%2==0 else BEGE)
    tb,tf=box(s,x+0.3,3.8,1.9,3.1)
    run(tf.paragraphs[0],f"{i+1:02d}",27,AREIA,F_BODY,bold=True)
    run(tf.add_paragraph(),t,19,PRETO,F_BODY,bold=True)
    if i<5: txt(s,"→",x+2.44,4.9,0.5,0.6,21,AREIA,bold=True)
rect(s,1.7,7.8,16.6,0.03,AREIA)
tb,tf=box(s,1.7,8.25,16.6,1.2); p=tf.paragraphs[0]
run(p,"Enxergar além   →   colocar em movimento   →   ",30,CORPO,F_BODY)
run(p,"fazer existir",30,PRETO,F_BODY,bold=True)

# =====================================================================
# 17 · A LÓGICA QUE NÃO MUDA  (manifesto)
# =====================================================================
s=add_slide()
eyebrow(s,"Os padrões",2.3,1.6)
txt(s,"Ao longo da sua trajetória, a Flávia repete um mesmo movimento:",2.3,2.25,14.6,0.8,21,CORPO)
tb,tf=box(s,2.3,3.1,15.4,3.4); p=tf.paragraphs[0]
run(p,"A Flávia ",56,PRETO,F_BODY)
run(p,"não aceita",56,PRETO,F_BODY,bold=True)
run(p," que aquilo que existe, determine aquilo que ",56,PRETO,F_BODY)
run(p,"pode existir.",56,PRETO,F_BODY,bold=True)
rect(s,2.5,6.85,15,0.04,AREIA)
txt(s,"Essa lógica aparece em contextos completamente diferentes da sua vida.",2.5,7.2,15,0.7,21,CORPO)
for i,t in enumerate(["Muda a situação.","Muda o desafio.","Muda o que precisa ser construído."]):
    x=2.5+i*5.1
    rect(s,x,7.95,4.7,1.1,BEGE)
    txt(s,t,x+0.35,8.2,4.2,0.7,20,PRETO,bold=True)
txt(s,"A Flávia vê aquilo que existe, mas também enxerga o que aquilo ainda pode se tornar.",
    2.5,9.4,15,1.0,26,PRETO,bold=True)

# =====================================================================
# 18 · DIVISOR 04 — O IKIGAI  (cinza)
# =====================================================================
s=add_slide(grey=True); divisor(s,"04","O IKIGAI")

# =====================================================================
# 19 · O QUE MOVE A FLÁVIA?
# =====================================================================
s=add_slide()
eyebrow(s,"O Ikigai")
h1(s,"O QUE MOVE A FLÁVIA?")
txt(s,"Até aqui, a história revelou como a Flávia responde à realidade e entra em movimento. O Ikigai acrescenta outra dimensão: o que faz esse movimento ter sentido para ela.",
    1.7,2.9,14.6,1.6,23,CORPO)
perguntas=["O que a Flávia ama?","No que reconhece as suas forças?",
           "Onde encontra realização?","Como deseja contribuir para outras pessoas?"]
for i,q in enumerate(perguntas):
    x=1.7+(i%2)*8.8; y=4.7+(i//2)*2.2
    rect(s,x,y,8.0,1.9,BRANCO if i%2==0 else BEGE)
    txt(s,f"{i+1:02d}",x+0.45,y+0.4,0.9,0.9,24,AREIA,bold=True)
    txt(s,q,x+1.5,y+0.45,6.2,1.1,24,PRETO,bold=True)
nota_rodape(s,"Não buscamos apenas aquilo que ela gosta de fazer, e sim o que faz uma realização ter significado para a Flávia.",9.4,20)

# =====================================================================
# 20 · O MAPA DO IKIGAI
# =====================================================================
s=add_slide()
eyebrow(s,"O Ikigai")
h1(s,"O MAPA DO IKIGAI")
ikigai=[("O que a Flávia ama","A família. As pessoas. Os momentos de qualidade. As conversas e as trocas verdadeiras. As conexões.",BRANCO),
        ("No que a Flávia é boa","No conhecimento que construiu. Na persuasão. Na seriedade. Na persistência. No domínio daquilo que vende. Na capacidade de envolver pessoas.",BEGE),
        ("Como a Flávia gosta de contribuir","Sendo útil. Compartilhando conhecimentos e experiências. Ajudando pessoas e empresários. Criando oportunidades. Fazendo diferença na vida das pessoas.",BEGE),
        ("Onde a Flávia encontra realização","Ao ver algo ganhar forma. Ao transformar matéria em algo de valor. Ao perceber a alegria do cliente. Ao ver o sonho de outra pessoa se tornar concreto. Ao saber que aquilo que construiu também ampliou as possibilidades de alguém.",BRANCO)]
for i,(tit,desc,cor) in enumerate(ikigai):
    x=1.7+(i%2)*8.8; y=3.2+(i//2)*3.35
    rect(s,x,y,8.0,3.1,cor)
    txt(s,tit,x+0.5,y+0.35,7.0,0.9,24,PRETO,bold=True)
    txt(s,desc,x+0.5,y+1.25,7.0,1.7,19,CORPO)

# =====================================================================
# 21 · O CENTRO DO IKIGAI
# =====================================================================
s=add_slide()
eyebrow(s,"O Ikigai")
h1(s,"O CENTRO DO IKIGAI")
txt(s,"A Flávia não se realiza apenas conquistando para si. Ela encontra realização quando vê uma possibilidade se tornar concreta e produzir algo na vida de outras pessoas.",
    1.7,2.85,10.2,2.2,27,PRETO,bold=True)
for i,t in enumerate(["Um cliente realiza um sonho.","Uma equipe cresce.",
                      "Uma ideia sai do papel.","Uma pessoa recebe uma oportunidade."]):
    y=5.3+i*1.15
    rect(s,1.7,y,10.2,0.98,BEGE if i%2 else BRANCO)
    txt(s,t,2.15,y+0.22,9.4,0.7,21,PRETO,bold=True)
rect(s,12.7,2.85,5.6,5.6,SALVIA)
txt(s,"Sobre o legado que deseja deixar, a própria Flávia resume:",13.15,3.2,4.8,1.2,18,CORPO)
txt(s,"“",13.15,4.2,1.5,1.4,56,BRANCO,bold=True)
txt(s,"Fazer diferença na vida das pessoas.",13.15,5.4,4.8,2.4,28,PRETO,bold=True)
nota_rodape(s,"A Flávia também deseja criar uma fundação para oferecer uma profissão a crianças.",9.95,20)

# =====================================================================
# 22 · RAZÃO DE SER
# =====================================================================
s=add_slide()
eyebrow(s,"O Ikigai",2.3,2.2)
eyebrow(s,"Razão de ser",2.3,2.9,10,AREIA)
tb,tf=box(s,2.3,3.6,15.4,3.6); p=tf.paragraphs[0]
run(p,"Transformar possibilidades em ",56,PRETO,F_BODY)
run(p,"realizações concretas",56,PRETO,F_BODY,bold=True)
run(p," que também ampliem a vida de outras pessoas.",56,PRETO,F_BODY)
rect(s,2.5,7.9,15,0.04,AREIA)
txt(s,"A realização ganha sentido quando aquilo que a Flávia constrói também amplia possibilidades para outras pessoas.",
    2.5,8.3,15,1.2,23,CORPO)

# =====================================================================
# 23 · DIVISOR 05 — A ESSÊNCIA  (cinza)
# =====================================================================
s=add_slide(grey=True); divisor(s,"05","A ESSÊNCIA")

# =====================================================================
# 24 · ONDE A HISTÓRIA E O IKIGAI SE ENCONTRAM
# =====================================================================
s=add_slide()
eyebrow(s,"A essência")
h1(s,"ONDE A HISTÓRIA E O IKIGAI SE ENCONTRAM")
txt(s,"A história revela como a Flávia se movimenta. O Ikigai revela o que dá sentido a esse movimento. Lado a lado, a mesma lógica aparece.",
    1.7,2.85,14.6,1.4,23,CORPO)
encontro=[("Visão","A Flávia enxerga além da condição presente."),
          ("Movimento","A distância entre aquilo que existe e aquilo que ela enxerga como possível a leva a buscar caminhos, conhecimento, pessoas e recursos."),
          ("Pessoas","A família, a equipe, os clientes e as relações fazem parte daquilo que dá significado às suas realizações."),
          ("Impacto","A conquista ganha mais sentido quando aquilo que ela constrói também amplia possibilidades para outras pessoas.")]
for i,(tit,desc) in enumerate(encontro):
    x=1.7+i*4.23
    rect(s,x,4.4,3.95,4.0,BRANCO if i%2==0 else BEGE)
    tb,tf=box(s,x+0.4,4.7,3.2,3.5)
    run(tf.paragraphs[0],tit,26,PRETO,F_BODY,bold=True)
    run(tf.add_paragraph(),desc,18,CORPO,F_BODY)
tb,tf=box(s,1.7,8.9,16.6,1.0); p=tf.paragraphs[0]
run(p,"Enxergar além.   ·   Colocar em movimento.   ·   ",26,CORPO,F_BODY)
run(p,"Fazer existir.",26,PRETO,F_BODY,bold=True)

# =====================================================================
# 25 · A ESSÊNCIA DA FLÁVIA COMO FUNDADORA
# =====================================================================
s=add_slide()
eyebrow(s,"A essência da Flávia como fundadora",1.7,1.5)
tb,tf=box(s,1.7,2.2,16.0,2.8); p=tf.paragraphs[0]
run(p,"Enxergar além do que está posto e ",52,PRETO,F_BODY)
run(p,"fazer existir",52,PRETO,F_BODY,bold=True)
run(p," o que ainda é possibilidade.",52,PRETO,F_BODY)
pilares=[("Enxergar além","A Flávia não considera a realidade presente como a única possibilidade. Ela consegue enxergar aquilo que ainda pode existir.",BRANCO),
         ("Fazer existir","Ela não permanece apenas no campo da imaginação. Procura caminhos e transforma a possibilidade em realidade.",BEGE),
         ("Ampliar possibilidades","A realização ganha sentido quando aquilo que ela constrói também cria valor, oportunidade ou transformação para outras pessoas.",SALVIA)]
for i,(tit,desc,cor) in enumerate(pilares):
    x=1.7+i*5.65
    rect(s,x,5.3,5.3,3.5,cor)
    eyebrow(s,tit,x+0.45,5.6,4.6)
    txt(s,desc,x+0.45,6.2,4.5,2.4,19,PRETO)
nota_rodape(s,"Forte, determinada, criativa, sonhadora e visionária são características da Flávia. Mas são manifestações de uma lógica mais profunda: enxergar além e fazer existir.",9.3,20)

# =====================================================================
# 26 · DIVISOR 06 — AS MARCAS  (cinza)
# =====================================================================
s=add_slide(grey=True); divisor(s,"06","AS MARCAS")

# =====================================================================
# 27 · COMO ESSA ESSÊNCIA SE MANIFESTA NAS MARCAS
# =====================================================================
s=add_slide()
eyebrow(s,"As marcas")
h1(s,"COMO ESSA ESSÊNCIA SE MANIFESTA NAS MARCAS")
txt(s,"Uma mesma essência. Duas manifestações diferentes.",1.7,2.8,14.6,0.9,27,PRETO,bold=True)
marcas=[("A Forma","Dar forma ao que ainda é ideia.",
         "A empresa representa o espaço que a Flávia buscava para criar, personalizar e construir segundo a sua própria visão. Esse mesmo movimento está na natureza do negócio: uma necessidade, uma ideia ou um sonho se transformam em projeto e ganham forma concreta.",BRANCO),
        ("A My Home","Reabrir possibilidades.",
         "A marca nasce quando a Flávia questiona o destino dado às sobras da marcenaria. Onde havia descarte, ela enxergou matéria. Onde havia fim, ela enxergou um novo começo. Aquilo que havia encerrado a sua função ganhou uma nova possibilidade de existir.",SALVIA)]
for i,(rot,frase,desc,cor) in enumerate(marcas):
    x=1.7+i*8.8
    rect(s,x,4.0,8.0,5.2,cor)
    eyebrow(s,rot,x+0.5,4.35,6.8)
    txt(s,frase,x+0.5,4.95,7.0,1.2,28,PRETO,bold=True)
    txt(s,desc,x+0.5,6.2,7.0,2.8,19,CORPO)
nota_rodape(s,"A Forma e a My Home são marcas diferentes e terão estratégias próprias. Mas as duas carregam, na origem, uma mesma forma de enxergar e transformar a realidade.",9.55,20)

# =====================================================================
# 28 · FRASE-MANIFESTO DE TRANSIÇÃO
# =====================================================================
s=add_slide()
tb,tf=box(s,2.3,4.0,15.4,3.4)
p=tf.paragraphs[0]
run(p,"A essência revela a ",62,PRETO,F_BODY); run(p,"origem.",62,PRETO,F_BODY,bold=True)
p2=tf.add_paragraph()
run(p2,"A estratégia define a ",62,PRETO,F_BODY); run(p2,"direção.",62,PRETO,F_BODY,bold=True)
rect(s,2.5,7.7,15,0.04,AREIA)

# =====================================================================
# 29 · DA ESSÊNCIA À ESTRATÉGIA
# =====================================================================
s=add_slide()
rect(s,11.4,0,8.6,11.25,BEGE2)
eyebrow(s,"Próxima etapa",1.5,1.0,9.0)
h1(s,"DA ESSÊNCIA À ESTRATÉGIA",1.5,1.6,39,9.4)
txt(s,"Esta etapa revelou a essência da Flávia como fundadora: uma forma própria de enxergar e transformar a realidade que já existia antes da Forma e da My Home.",
    1.5,3.4,9.2,2.6,24,CORPO)
eyebrow(s,"O que vem a seguir",12.0,2.4,7.4)
txt(s,"Essa essência ajuda a compreender de onde as duas marcas vêm, mas não define, sozinha, quem cada marca precisa ser.",
    12.0,3.1,7.2,2.2,23,CORPO)
rect(s,12.0,5.6,7.2,0.03,AREIA)
txt(s,"A Base Estratégica vai definir como essa origem se traduz em uma direção própria, relevante e diferenciada para a Forma e para a My Home.",
    12.0,6.0,7.2,3.0,26,PRETO,bold=True)

# =====================================================================
# 30 · ENCERRAMENTO
# =====================================================================
s=add_slide()
tb,tf=box(s,1.7,4.0,16,2); run(tf.paragraphs[0],"MUITO OBRIGADA!",70,PRETO0,F_TIT)
txt(s,"Aplicando os conteúdos deste documento na prática, você garante uma marca construída com consistência e clareza.",
    1.7,6.2,14,1.6,30,PRETO)

out="/home/user/apresentacoes/REVELACAO-DE-ESSENCIA-FLAVIA-PEREIRA-MUCCELIN.pptx"
prs.save(out)
print("salvo:",out,"| slides:",len(prs.slides._sldIdLst))
